"""Render a tiny deck with the rule-based drafter and re-parse it."""

from __future__ import annotations

import tempfile
from pathlib import Path

from pptx import Presentation

from pt_maker.render.engine import render_deck
from pt_maker.slides.drafter import draft_slides_rule
from pt_maker.theme.schema import default_theme


def test_render_round_trip():
    theme = default_theme("t")
    outline = {
        "meta": {"title": "Demo", "purpose": "business", "language": "ko"},
        "slides": [
            {"id": "s1", "type": "title", "title": "데모 프레젠테이션", "subtitle": "PT_maker 검증"},
            {"id": "s2", "type": "content", "title": "개요",
             "bullets": ["첫 번째 항목", "두 번째 항목", "세 번째 항목"]},
            {"id": "s3", "type": "chart", "title": "연도별 매출",
             "chart": {"kind": "bar",
                       "data": {"categories": ["2022", "2023", "2024"],
                                "series": [{"name": "매출", "values": [10, 14, 18]}]}}},
            {"id": "s4", "type": "table", "title": "KPI",
             "headers": ["지표", "값"], "rows": [["MAU", "1.2M"], ["NPS", "42"]]},
            {"id": "s5", "type": "closing", "title": "감사합니다"},
        ],
    }
    spec = draft_slides_rule(outline, theme)
    assert len(spec["slides"]) == len(outline["slides"])

    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "test.pptx"
        render_deck(spec, theme, out)
        assert out.exists()
        prs = Presentation(str(out))
        assert len(prs.slides) == len(outline["slides"])

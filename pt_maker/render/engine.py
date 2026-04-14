"""Render a slides spec + theme into a ``.pptx`` file.

Design choice
-------------
We do **not** rewrite the OOXML slide master. Instead, we use a blank
presentation sized to the theme and draw every element directly with absolute
positioning derived from the theme's layout definitions. This keeps the
renderer deterministic and cross-engine (LibreOffice) robust.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches

from ..components import chart as chart_c
from ..components import diagram as diagram_c
from ..components import icon as icon_c
from ..components import image as image_c
from ..components import math as math_c
from ..components import shape as shape_c
from ..components import table as table_c
from ..components import text as text_c
from .fonts import resolve_color


def _apply_background(slide, layout: dict[str, Any], theme: dict[str, Any]) -> None:
    token = layout.get("background_token", "background")
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = resolve_color(token, theme)


def _expand_placeholder(block: dict[str, Any], layout: dict[str, Any]) -> dict[str, Any]:
    """If block uses ``placeholder_role`` without explicit geometry, copy from layout."""
    if "x_in" in block and "y_in" in block and "w_in" in block and "h_in" in block:
        return block
    role = block.get("placeholder_role")
    if not role:
        return block
    for ph in layout.get("placeholders", []):
        if ph.get("role") == role:
            merged = dict(ph)
            merged.update(block)
            return merged
    return block


def _add_decorations(slide, theme: dict[str, Any], layout_name: str) -> None:
    dec = theme.get("decorations", {})
    bar = dec.get("header_bar", {})
    if bar.get("enabled") and layout_name not in ("title", "section_divider", "closing"):
        shape_c.add_shape(slide, {
            "shape": "RECTANGLE",
            "x_in": 0, "y_in": 0,
            "w_in": theme["slide_size"]["width_in"],
            "h_in": float(bar.get("height_in", 0.3)),
            "fill_token": bar.get("fill_token", "primary"),
        }, theme)


def render_deck(spec: dict[str, Any], theme: dict[str, Any], out_path: str | Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(theme["slide_size"]["width_in"])
    prs.slide_height = Inches(theme["slide_size"]["height_in"])

    blank_layout = prs.slide_layouts[6]  # blank

    for page_num, s in enumerate(spec["slides"], start=1):
        slide = prs.slides.add_slide(blank_layout)
        layout_name = s.get("layout", "content")
        layout = theme["layouts"].get(layout_name) or theme["layouts"]["content"]

        _apply_background(slide, layout, theme)
        _add_decorations(slide, theme, layout_name)

        for block in s.get("blocks", []):
            block = _expand_placeholder(block, layout)
            kind = block.get("kind")
            try:
                if kind == "text":
                    # Title-like when role is title
                    if block.get("placeholder_role") in ("title", "subtitle"):
                        text_c.add_title(slide, block, theme)
                    else:
                        text_c.add_text_block(slide, block, theme)
                elif kind == "bullets":
                    text_c.add_bullets(slide, block, theme)
                elif kind == "image":
                    image_c.add_image(slide, block, theme)
                elif kind == "chart":
                    chart_c.add_chart(slide, block, theme)
                elif kind == "table":
                    table_c.add_table(slide, block, theme)
                elif kind == "math":
                    math_c.add_math(slide, block, theme)
                elif kind == "icon":
                    icon_c.add_icon(slide, block, theme)
                elif kind == "diagram":
                    diagram_c.add_diagram(slide, block, theme)
                elif kind == "shape":
                    shape_c.add_shape(slide, block, theme)
            except Exception as e:  # keep rendering, report per-slide
                _add_error_note(slide, f"[render error: {kind}: {e}]")

        # Speaker notes
        notes = s.get("notes")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

        # Footer / page number
        dec = theme.get("decorations", {}).get("footer", {})
        if dec.get("enabled", True) and layout_name not in ("title", "section_divider"):
            _add_footer(slide, theme, dec, page_num=page_num)

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


def _add_footer(slide, theme, dec, page_num: int) -> None:
    from .fonts import apply_run_fonts

    w_in = theme["slide_size"]["width_in"]
    h_in = theme["slide_size"]["height_in"]
    typo = theme["typography"]
    size = typo["sizes_pt"].get("caption", 11)

    txt = dec.get("text", "")
    if dec.get("page_number", True):
        txt = f"{txt}   |   {page_num}" if txt else f"{page_num}"
    if not txt.strip():
        return

    tb = slide.shapes.add_textbox(
        Inches(0.5), Inches(h_in - 0.4), Inches(w_in - 1.0), Inches(0.3),
    )
    tf = tb.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = txt
    apply_run_fonts(
        run,
        latin=typo.get("body_font_latin", "Inter"),
        ea=typo.get("body_font_ea", "Pretendard"),
        size_pt=size,
        color=resolve_color("muted", theme),
    )


def _add_error_note(slide, msg: str) -> None:
    tb = slide.shapes.add_textbox(Inches(0.2), Inches(0.2), Inches(6), Inches(0.3))
    tb.text_frame.text = msg

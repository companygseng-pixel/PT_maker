"""Text + bullet list components."""

from __future__ import annotations

from typing import Any

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from ..render.fonts import apply_run_fonts, resolve_color


_ALIGN_MAP = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}


def add_text_block(slide, block: dict[str, Any], theme: dict[str, Any]) -> None:
    text = block.get("text", "")
    size = float(block.get("font_size_pt", theme["typography"]["sizes_pt"].get("body", 18)))
    color_token = block.get("color_token", "text")
    align = _ALIGN_MAP.get(block.get("align", "left"), PP_ALIGN.LEFT)
    bold = bool(block.get("bold", False))

    tb = slide.shapes.add_textbox(
        Inches(block["x_in"]), Inches(block["y_in"]),
        Inches(block["w_in"]), Inches(block["h_in"]),
    )
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text

    typo = theme["typography"]
    apply_run_fonts(
        run,
        latin=typo.get("body_font_latin", "Inter"),
        ea=typo.get("body_font_ea", "Pretendard"),
        size_pt=size,
        color=resolve_color(color_token, theme),
        bold=bold,
    )
    p.line_spacing = 1.25


def add_bullets(slide, block: dict[str, Any], theme: dict[str, Any]) -> None:
    items: list[str] = block.get("items", [])
    size = float(block.get("font_size_pt", theme["typography"]["sizes_pt"].get("body", 18)))
    color_token = block.get("color_token", "text")

    tb = slide.shapes.add_textbox(
        Inches(block["x_in"]), Inches(block["y_in"]),
        Inches(block["w_in"]), Inches(block["h_in"]),
    )
    tf = tb.text_frame
    tf.word_wrap = True

    typo = theme["typography"]
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = 1.3
        p.space_after = Pt(4)
        p.level = 0
        run = p.add_run()
        run.text = f"•  {item}"
        apply_run_fonts(
            run,
            latin=typo.get("body_font_latin", "Inter"),
            ea=typo.get("body_font_ea", "Pretendard"),
            size_pt=size,
            color=resolve_color(color_token, theme),
        )


def add_title(slide, block: dict[str, Any], theme: dict[str, Any]) -> None:
    text = block.get("text", "")
    size = float(block.get("font_size_pt", theme["typography"]["sizes_pt"].get("heading", 28)))
    color_token = block.get("color_token", "primary")
    align = _ALIGN_MAP.get(block.get("align", "left"), PP_ALIGN.LEFT)

    tb = slide.shapes.add_textbox(
        Inches(block["x_in"]), Inches(block["y_in"]),
        Inches(block["w_in"]), Inches(block["h_in"]),
    )
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    typo = theme["typography"]
    apply_run_fonts(
        run,
        latin=typo.get("title_font_latin", typo.get("body_font_latin", "Inter")),
        ea=typo.get("title_font_ea", typo.get("body_font_ea", "Pretendard")),
        size_pt=size,
        color=resolve_color(color_token, theme),
        bold=True,
    )

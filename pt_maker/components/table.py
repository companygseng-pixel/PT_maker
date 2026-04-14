"""Styled tables."""

from __future__ import annotations

from typing import Any

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from ..render.fonts import apply_run_fonts, resolve_color


def add_table(slide, block: dict[str, Any], theme: dict[str, Any]) -> None:
    rows = block.get("rows") or []
    headers = block.get("headers") or (rows[0] if rows else [])
    body_rows = rows[1:] if (rows and rows[0] is headers) else rows

    n_rows = len(body_rows) + (1 if headers else 0)
    n_cols = max(len(r) for r in rows) if rows else len(headers)
    if n_rows == 0 or n_cols == 0:
        return

    shape = slide.shapes.add_table(
        n_rows, n_cols,
        Inches(block["x_in"]), Inches(block["y_in"]),
        Inches(block["w_in"]), Inches(block["h_in"]),
    )
    table = shape.table
    typo = theme["typography"]
    body_size = float(block.get("font_size_pt", typo["sizes_pt"].get("body", 14)))

    primary = resolve_color("primary", theme)
    surface = resolve_color("surface", theme)
    text_color = resolve_color("text", theme)
    bg_color = resolve_color("background", theme)

    row_idx = 0
    if headers:
        for ci, h in enumerate(headers):
            cell = table.cell(0, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = primary
            tf = cell.text_frame
            tf.text = ""
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = str(h)
            apply_run_fonts(
                run,
                latin=typo.get("title_font_latin", "Inter"),
                ea=typo.get("title_font_ea", "Pretendard"),
                size_pt=body_size, color=bg_color, bold=True,
            )
        row_idx = 1

    for ri, row in enumerate(body_rows):
        r = row_idx + ri
        for ci in range(n_cols):
            cell = table.cell(r, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = surface if ri % 2 else bg_color
            tf = cell.text_frame
            tf.text = ""
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = str(row[ci]) if ci < len(row) else ""
            apply_run_fonts(
                run,
                latin=typo.get("body_font_latin", "Inter"),
                ea=typo.get("body_font_ea", "Pretendard"),
                size_pt=body_size, color=text_color,
            )
